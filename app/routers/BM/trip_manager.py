"""后台管理 —— 出差记录查询 / 汇总 / 导出

查询口径：
- 一律按 cycle_key（结算月）过滤，边界日 25 号不会被相邻两个月重复统计；
- 明细支持按姓名、出差原因、具体日期再过滤；
- 汇总给出「每个人填了几条、几天 + 具体日期 + 最后提交时间」，并列出周期内一条都没填的账号。
"""

import tempfile
import urllib.parse
from datetime import datetime
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db_services.database import get_db
from app.dependencies.BM_auth import bm_verify_token
from app.logger import get_logger
from app.models.trip import TripRecord
from app.models.user import User
from app.utils.enveloping_route import EnvelopingRoute
from app.utils.trip_cycle import cycle_by_key, current_cycle, recent_cycle_keys, to_cn

router = APIRouter(route_class=EnvelopingRoute)
logger = get_logger("trip_manager")

WEEKDAY_CN = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]


def _resolve_cycle(cycle: Optional[str]) -> tuple:
    try:
        return cycle_by_key(cycle) if cycle else current_cycle()
    except ValueError:
        raise HTTPException(status_code=400, detail="结算月格式应为 YYYY-MM")


@router.get("/cycles", summary="可查询的结算月列表")
async def list_cycles(token_payload: dict = Depends(bm_verify_token)):
    start, end, key = current_cycle()
    options = []
    for k in recent_cycle_keys(12):
        s, e, _ = cycle_by_key(k)
        options.append(
            {
                "cycle_key": k,
                "start_date": s.isoformat(),
                "end_date": e.isoformat(),
                "label": f"{k}（{s:%m-%d} ~ {e:%m-%d}）",
                "is_current": k == key,
            }
        )
    return {"current": key, "current_start": start, "current_end": end, "options": options}


@router.get("/list", summary="出差记录明细（分页）")
async def list_records(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    cycle: Optional[str] = Query(None, description="结算月 YYYY-MM"),
    name: Optional[str] = Query(None, description="姓名模糊匹配"),
    reason: Optional[str] = Query(None, description="出差原因模糊匹配"),
    trip_date: Optional[str] = Query(None, description="精确到某一天 YYYY-MM-DD"),
    db: AsyncSession = Depends(get_db),
    token_payload: dict = Depends(bm_verify_token),
):
    start, end, key = _resolve_cycle(cycle)
    try:
        # 姓名不落库，统一 join user 表取最新姓名
        stmt = (
            select(TripRecord, User.name)
            .outerjoin(User, User.id == TripRecord.user_id)
            .where(TripRecord.cycle_key == key)
        )
        if name:
            stmt = stmt.where(User.name.ilike(f"%{name}%"))
        if reason:
            stmt = stmt.where(TripRecord.trip_reason.ilike(f"%{reason}%"))
        if trip_date:
            try:
                exact = datetime.strptime(trip_date, "%Y-%m-%d").date()
            except ValueError:
                raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
            stmt = stmt.where(TripRecord.trip_date == exact)

        rows = (
            await db.execute(
                stmt.order_by(
                    TripRecord.trip_date.asc(),
                    User.name.asc(),
                    TripRecord.id.asc(),
                )
            )
        ).all()
        total = len(rows)
        page_rows = rows[(page - 1) * page_size: page * page_size]

        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "cycle_key": key,
            "start_date": start.isoformat(),
            "end_date": end.isoformat(),
            "items": [
                {
                    "id": r.id,
                    "user_id": r.user_id,
                    "user_name": user_name or "（账号已删除）",
                    "trip_date": r.trip_date.isoformat(),
                    "weekday": WEEKDAY_CN[r.trip_date.weekday()],
                    "trip_reason": r.trip_reason,
                    "cycle_key": r.cycle_key,
                    "create_at": to_cn(r.create_at),
                }
                for r, user_name in page_rows
            ],
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"查询出差明细失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail="查询出差明细失败")


@router.get("/summary", summary="按人汇总填写情况")
async def summary_by_user(
    cycle: Optional[str] = Query(None, description="结算月 YYYY-MM"),
    db: AsyncSession = Depends(get_db),
    token_payload: dict = Depends(bm_verify_token),
):
    start, end, key = _resolve_cycle(cycle)
    try:
        rows = (
            await db.execute(
                select(TripRecord)
                .where(TripRecord.cycle_key == key)
                .order_by(TripRecord.trip_date.asc(), TripRecord.id.asc())
            )
        ).scalars().all()

        # 姓名一次查出来做映射，避免 N+1
        all_users = (await db.execute(select(User.id, User.name, User.is_active, User.user_type))).all()
        name_map = {uid: name for uid, name, _active, _t in all_users}

        grouped: dict = {}
        for r in rows:
            item = grouped.setdefault(
                r.user_id,
                {
                    "user_id": r.user_id,
                    "dates": set(),
                    "records": 0,
                    "last_submit": None,
                },
            )
            item["dates"].add(r.trip_date)
            item["records"] += 1
            if r.create_at and (item["last_submit"] is None or r.create_at > item["last_submit"]):
                item["last_submit"] = r.create_at

        items = []
        for item in grouped.values():
            items.append(
                {
                    "user_id": item["user_id"],
                    "user_name": name_map.get(item["user_id"]) or "（账号已删除）",
                    "filled_days": len(item["dates"]),
                    "record_count": item["records"],
                    "dates": sorted(d.isoformat() for d in item["dates"]),
                    "last_submit": to_cn(item["last_submit"]),
                }
            )
        items.sort(key=lambda x: (-x["record_count"], x["user_name"]))

        # 周期内一条没填的启用账号（只统计普通账号，管理员不参与填报）
        filled_ids = set(grouped.keys())
        unfilled = sorted(
            name for uid, name, active, utype in all_users
            if active and utype != "admin" and uid not in filled_ids
        )

        return {
            "cycle_key": key,
            "start_date": start.isoformat(),
            "end_date": end.isoformat(),
            "total_records": len(rows),
            "total_users": len(grouped),
            "unfilled_users": unfilled,
            "items": items,
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"查询出差汇总失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail="查询出差汇总失败")


@router.get("/export/excel", summary="按结算月导出 Excel")
async def export_excel(
    cycle: Optional[str] = Query(None, description="结算月 YYYY-MM"),
    name: Optional[str] = Query(None, description="仅导出指定姓名（模糊）"),
    reason: Optional[str] = Query(None, description="仅导出指定出差原因（模糊）"),
    db: AsyncSession = Depends(get_db),
    token_payload: dict = Depends(bm_verify_token),
):
    start, end, key = _resolve_cycle(cycle)
    logger.info(f"导出出差记录 Excel cycle={key} name={name}")

    try:
        stmt = (
            select(TripRecord, User.name)
            .outerjoin(User, User.id == TripRecord.user_id)
            .where(TripRecord.cycle_key == key)
        )
        if name:
            stmt = stmt.where(User.name.ilike(f"%{name}%"))
        if reason:
            stmt = stmt.where(TripRecord.trip_reason.ilike(f"%{reason}%"))
        rows = (
            await db.execute(
                stmt.order_by(
                    TripRecord.trip_date.asc(),
                    User.name.asc(),
                    TripRecord.id.asc(),
                )
            )
        ).all()

        wb = openpyxl.Workbook()

        # ---------- Sheet1 明细 ----------
        ws = wb.active
        ws.title = "出差明细"
        headers = ["序号", "姓名", "出差日期", "星期", "出差原因", "提交时间", "结算月"]
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        for col, title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, (r, user_name) in enumerate(rows, 1):
            ws.append(
                [
                    idx,
                    user_name or "（账号已删除）",
                    r.trip_date.isoformat(),
                    WEEKDAY_CN[r.trip_date.weekday()],
                    r.trip_reason,
                    to_cn(r.create_at).strftime("%Y-%m-%d %H:%M:%S") if r.create_at else "",
                    r.cycle_key,
                ]
            )

        for col, width in enumerate([6, 12, 14, 8, 60, 20, 10], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
        # 出差原因是自由文本，可能带换行，纵向靠上并自动换行
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"

        # ---------- Sheet2 人员汇总 ----------
        ws2 = wb.create_sheet("人员汇总")
        headers2 = ["姓名", "出差天数", "已填日期", "最后提交时间"]
        for col, title in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        grouped: dict = {}
        for r, user_name in rows:
            item = grouped.setdefault(
                r.user_id, {"name": user_name or "（账号已删除）", "dates": set(), "records": 0, "last": None}
            )
            item["dates"].add(r.trip_date)
            item["records"] += 1
            if r.create_at and (item["last"] is None or r.create_at > item["last"]):
                item["last"] = r.create_at

        for item in sorted(grouped.values(), key=lambda x: (-x["records"], x["name"])):
            ws2.append(
                [
                    item["name"],
                    len(item["dates"]),
                    "、".join(sorted(d.isoformat() for d in item["dates"])),
                    to_cn(item["last"]).strftime("%Y-%m-%d %H:%M:%S") if item["last"] else "",
                ]
            )

        for col, width in enumerate([12, 12, 60, 20], 1):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
        ws2.freeze_panes = "A2"

        filename = f"出差记录_{key}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            wb.save(tmp.name)
            temp_path = tmp.name

        encoded = urllib.parse.quote(filename.encode("utf-8"))
        return FileResponse(
            path=temp_path,
            filename=filename,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"导出出差记录失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")
